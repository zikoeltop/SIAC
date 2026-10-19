"""Project site visits and periodic asset inspections.

All values recorded here are visit snapshots. Editing an inspection never writes to Asset.
"""
from collections import defaultdict
from datetime import date
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Count, Q
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from audit.utils import log_event
from .models import Asset, Employee, SiteInspection, SiteVisit


PHYSICAL_TYPES = [Asset.Types.COMPUTER, Asset.Types.MONITOR, Asset.Types.PRINTER, Asset.Types.OTHER]
EXCLUDED_STATUSES = [Asset.Status.SCRAPPED, Asset.Status.ARCHIVED]
SUCCESS_STATUSES = [SiteInspection.Status.INSPECTED, SiteInspection.Status.NEEDS_ATTENTION]
EDITABLE_SNAPSHOT_FIELDS = {
    'user_name': ('USER NAME', 'سبب تغيير اسم المستخدم'),
    'monitor_model': ('Monitor Model', 'سبب تغيير موديل الشاشة'),
    'monitor_serial': ('Monitor Serial', 'سبب تغيير سيريال الشاشة'),
    'printer_model': ('printer model', 'سبب تغيير موديل الطابعة'),
    'printer_serial': ('printer serial', 'سبب تغيير سيريال الطابعة'),
}
EXCEL_HEADERS = [
    'Project', 'visit date', 'IT TECH NAME', 'USER NAME', 'TYPE', 'MODEL',
    'SERIAL NUMBER', 'COMPUTER NAME', 'Monitor Model', 'Monitor Serial',
    'printer model', 'printer serial', 'COMMENTS',
]


def _inspectable_assets():
    return Asset.objects.filter(asset_type__in=PHYSICAL_TYPES).exclude(asset_status__in=EXCLUDED_STATUSES)


def _project_scoped_assets(project):
    return _inspectable_assets().filter(
        Q(project__iexact=project) | (Q(project='') & Q(employee__project__iexact=project))
    )


def _asset_belongs_to_project(asset, project):
    asset_project = (asset.project or '').strip()
    if asset_project:
        return asset_project.casefold() == (project or '').strip().casefold()
    employee_project = (asset.employee.project or '').strip() if asset.employee_id else ''
    return employee_project.casefold() == (project or '').strip().casefold()


def _project_names():
    """Unique normalized project labels based on assets, plus employee project references."""
    values = list(_inspectable_assets().exclude(project='').values_list('project', flat=True))
    values += list(Employee.objects.exclude(project='').values_list('project', flat=True))
    unique = {}
    for value in values:
        value = (value or '').strip()
        if value:
            unique.setdefault(value.casefold(), value)
    return sorted(unique.values(), key=str.casefold)


def _type_snapshot(asset):
    if asset.asset_type == Asset.Types.COMPUTER:
        type_text = ' '.join([(asset.item_type or ''), (asset.model or '')]).casefold()
        return 'LAPTOP' if ('laptop' in type_text or 'notebook' in type_text) else 'PC'
    if asset.asset_type == Asset.Types.MONITOR:
        return 'MONITOR'
    if asset.asset_type == Asset.Types.PRINTER:
        return 'PRINTER'
    return (asset.item_type or asset.get_asset_type_display() or 'OTHER').strip()[:120]


def _asset_display_name(asset):
    name = (asset.display_employee_name or '').strip()
    return '' if name == '-' else name


def _join_unique(values, max_length=255):
    result = []
    seen = set()
    for value in values:
        value = (value or '').strip()
        key = value.casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return '; '.join(result)[:max_length]


def _normalized_match(value):
    """Normalize names/IDs from imported inventory columns for reliable matching."""
    return ' '.join(str(value or '').strip().casefold().split())


def _asset_person_keys(asset):
    employee = getattr(asset, 'employee', None)
    values = [
        getattr(asset, 'hr_code', ''),
        getattr(employee, 'hr_code', '') if employee else '',
        getattr(asset, 'employee_name', ''),
        getattr(employee, 'name', '') if employee else '',
        getattr(employee, 'arabic_name', '') if employee else '',
        getattr(asset, 'display_employee_name', ''),
    ]
    return {_normalized_match(value) for value in values if _normalized_match(value) and _normalized_match(value) != '-'}


def _source_identity_values(asset, max_columns):
    """Read identity fields only, while limiting the scan to A:K or A:I."""
    raw = asset.raw_data if isinstance(asset.raw_data, dict) else {}
    identity_terms = (
        'siac id', 'user name', 'username', 'hr id', 'hr code', 'employee',
        'computer name', 'computer', 'asset number', 'asset tag', 'serial',
        'user id', 'name',
    )
    values = []
    for header, value in list(raw.items())[:max_columns]:
        normalized_header = _normalized_match(header)
        if any(term in normalized_header for term in identity_terms):
            values.append(value)
    return values


def _peripheral_map(assets):
    """Match peripherals to computers using identifiers in the imported source rows.

    Monitors 2 is searched across A:K and Printers 2027 across A:I. Model/serial
    values are later paired from each matched peripheral row, and Asset records
    are never modified by visit inspection edits.
    """
    computers = [asset for asset in assets if asset.asset_type == Asset.Types.COMPUTER]
    result = defaultdict(lambda: {'monitors': [], 'printers': []})
    if not computers:
        return result

    aliases_by_computer = {}
    for computer in computers:
        aliases = _asset_person_keys(computer)
        for value in (computer.serial_number, computer.asset_tag, computer.computer_name, computer.computer_name_new):
            key = _normalized_match(value)
            if key and key != '-':
                aliases.add(key)
        aliases_by_computer[computer.pk] = aliases

    related = Asset.objects.filter(
        asset_type__in=[Asset.Types.MONITOR, Asset.Types.PRINTER]
    ).exclude(asset_status__in=EXCLUDED_STATUSES).select_related('employee').order_by('id')

    for peripheral in related:
        max_columns = 11 if peripheral.asset_type == Asset.Types.MONITOR else 9
        source_values = _source_identity_values(peripheral, max_columns)
        source_values.extend([
            peripheral.hr_code, peripheral.asset_tag, peripheral.employee_name,
            peripheral.employee.hr_code if peripheral.employee_id else '',
            peripheral.employee.name if peripheral.employee_id else '',
            peripheral.employee.arabic_name if peripheral.employee_id else '',
        ])
        row_values = {_normalized_match(value) for value in source_values if _normalized_match(value) and _normalized_match(value) != '-'}
        bucket = 'monitors' if peripheral.asset_type == Asset.Types.MONITOR else 'printers'
        for computer in computers:
            if aliases_by_computer[computer.pk].intersection(row_values):
                key = _normalized_match(computer.hr_code) or f'asset:{computer.pk}'
                result[key][bucket].append(peripheral)
    return result


def _paired_values(items, model_getter):
    """Return model and serial lists from the same deduplicated source rows."""
    pairs, seen = [], set()
    for item in items:
        model = str(model_getter(item) or '').strip()
        serial = str(item.serial_number or '').strip()
        pair = (model.casefold(), serial.casefold())
        if (model or serial) and pair not in seen:
            seen.add(pair)
            pairs.append((model or '-', serial or '-'))
    return (
        '; '.join(model for model, _ in pairs)[:255],
        '; '.join(serial for _, serial in pairs)[:255],
    )


def _snapshot_defaults(asset, visit, is_additional=False, peripherals=None):
    peripherals = peripherals or {}
    code = _normalized_match(asset.hr_code)
    # Peripheral fields are populated only on computer rows. Keep all changes in
    # inspection snapshots; never write edited values back to the source Asset.
    matching_key = code or f'asset:{asset.pk}'
    matching = peripherals.get(matching_key, {'monitors': [], 'printers': []}) if asset.asset_type == Asset.Types.COMPUTER else {'monitors': [], 'printers': []}
    monitors = matching.get('monitors', [])
    printers = matching.get('printers', [])

    monitor_models, monitor_serials = _paired_values(monitors, lambda item: item.model or item.item_type or item.brand)
    printer_models, printer_serials = _paired_values(printers, lambda item: item.model or item.item_type or item.brand)
    user_name = _asset_display_name(asset)
    return {
        'visit': visit,
        'asset': asset,
        'is_additional': is_additional,
        'source_project': (asset.project or '').strip() or ((asset.employee.project or '').strip() if asset.employee_id else ''),
        'source_user_name': user_name,
        'user_name': user_name,
        'type_snapshot': _type_snapshot(asset),
        'model_snapshot': (asset.model or asset.item_type or asset.brand or '').strip()[:255],
        'serial_number_snapshot': (asset.serial_number or '').strip()[:180],
        'computer_name_snapshot': (asset.computer_name or asset.computer_name_new or '').strip()[:255],
        'source_monitor_model': monitor_models,
        'monitor_model': monitor_models,
        'source_monitor_serial': monitor_serials,
        'monitor_serial': monitor_serials,
        'source_printer_model': printer_models,
        'printer_model': printer_models,
        'source_printer_serial': printer_serials,
        'printer_serial': printer_serials,
    }


def _refresh_peripheral_snapshots(visit):
    """Backfill peripheral snapshots for existing visits without overwriting edits."""
    computer_rows = list(visit.inspections.filter(asset__asset_type=Asset.Types.COMPUTER).select_related('asset', 'asset__employee'))
    if not computer_rows:
        return
    peripherals = _peripheral_map([row.asset for row in computer_rows])
    fields = (
        ('monitor_model', 'source_monitor_model'),
        ('monitor_serial', 'source_monitor_serial'),
        ('printer_model', 'source_printer_model'),
        ('printer_serial', 'source_printer_serial'),
    )
    for row in computer_rows:
        key = _normalized_match(row.asset.hr_code) or f'asset:{row.asset_id}'
        matched = peripherals.get(key, {'monitors': [], 'printers': []})
        defaults = {
            'source_monitor_model': _paired_values(matched['monitors'], lambda item: item.model or item.item_type or item.brand)[0],
            'source_monitor_serial': _paired_values(matched['monitors'], lambda item: item.model or item.item_type or item.brand)[1],
            'source_printer_model': _paired_values(matched['printers'], lambda item: item.model or item.item_type or item.brand)[0],
            'source_printer_serial': _paired_values(matched['printers'], lambda item: item.model or item.item_type or item.brand)[1],
        }
        changed_fields = []
        for current_field, source_field in fields:
            new_source = defaults[source_field]
            old_source = getattr(row, source_field) or ''
            current_value = getattr(row, current_field) or ''
            # Update empty/unmodified snapshots only. Technician edits remain untouched.
            if current_value == old_source or not current_value:
                if current_value != new_source:
                    setattr(row, current_field, new_source)
                    changed_fields.append(current_field)
            if old_source != new_source:
                setattr(row, source_field, new_source)
                changed_fields.append(source_field)
        if changed_fields:
            row.save(update_fields=list(dict.fromkeys(changed_fields + ['updated_at'])))


def _seed_project_assets(visit):
    """Ensure every current physical asset in the project's scope has a checklist row."""
    if visit.status == SiteVisit.Status.COMPLETED:
        return 0
    assets = list(_project_scoped_assets(visit.project).select_related('employee').order_by('id'))
    if not assets:
        return 0
    existing = set(visit.inspections.values_list('asset_id', flat=True))
    missing = [asset for asset in assets if asset.pk not in existing]
    if not missing:
        return 0
    peripherals = _peripheral_map(missing)
    rows = [SiteInspection(**_snapshot_defaults(asset, visit, False, peripherals)) for asset in missing]
    SiteInspection.objects.bulk_create(rows, ignore_conflicts=True)
    return len(rows)


def _create_additional_inspection(visit, asset):
    peripherals = _peripheral_map([asset])
    is_additional = not _asset_belongs_to_project(asset, visit.project)
    inspection = SiteInspection.objects.create(**_snapshot_defaults(asset, visit, is_additional, peripherals))
    if is_additional:
        source_project = inspection.source_project or '(blank / not assigned)'
        inspection.auto_comments = (
            f'[AUTO {timezone.localdate().isoformat()}] Added to visit for project "{visit.project}" as an additional on-site device. '
            f'SIAC source project: "{source_project}". The main SIAC asset record was not changed.'
        )
        inspection.save(update_fields=['auto_comments', 'updated_at'])
    return inspection


def _site_visit_aggregate(queryset):
    return queryset.annotate(
        inspection_total=Count('inspections', distinct=True),
        inspected_total=Count('inspections', filter=Q(inspections__status__in=SUCCESS_STATUSES), distinct=True),
        pending_total=Count('inspections', filter=Q(inspections__status=SiteInspection.Status.PENDING), distinct=True),
        unable_total=Count('inspections', filter=Q(inspections__status=SiteInspection.Status.UNABLE), distinct=True),
    )


def site_visit_global_metrics():
    """Shared global metrics for the root Smart Dashboard and management dashboard."""
    assets = list(_inspectable_assets().select_related('employee').only('id', 'project', 'employee__project'))
    asset_ids = [asset.id for asset in assets]
    successful_counts = dict(
        SiteInspection.objects.filter(asset_id__in=asset_ids, status__in=SUCCESS_STATUSES)
        .values('asset_id').annotate(total=Count('id')).values_list('asset_id', 'total')
    ) if asset_ids else {}
    latest_result_by_asset = {}
    recent_results = SiteInspection.objects.filter(
        asset_id__in=asset_ids
    ).exclude(status=SiteInspection.Status.PENDING).order_by('-visit__visit_date', '-inspected_at', '-id') if asset_ids else SiteInspection.objects.none()
    for inspection in recent_results.only('asset_id', 'status'):
        latest_result_by_asset.setdefault(inspection.asset_id, inspection.status)
    ever_ids = set(successful_counts)
    grouped_projects = defaultdict(list)
    for asset in assets:
        project = (asset.project or '').strip() or ((asset.employee.project or '').strip() if asset.employee_id else '')
        if project:
            grouped_projects.setdefault(project.casefold(), {'name': project, 'asset_ids': []})['asset_ids'].append(asset.id)

    latest_by_project = {}
    visits = SiteVisit.objects.all()
    for visit in visits.order_by('-visit_date', '-created_at').only('project', 'visit_date'):
        latest_by_project.setdefault(visit.project.casefold(), visit.visit_date)

    # Project coverage must only count inspections performed on that project's
    # own visits. A device inspected as an extra in another project must not
    # inflate the coverage percentage of its source project.
    project_inspected_assets = defaultdict(set)
    project_results = SiteInspection.objects.filter(
        status__in=SUCCESS_STATUSES, is_additional=False
    ).values('visit__project', 'asset_id').annotate(total=Count('id'))
    for item in project_results:
        key = (item['visit__project'] or '').strip().casefold()
        if key:
            project_inspected_assets[key].add(item['asset_id'])

    project_rows = []
    for value in grouped_projects.values():
        ids = value['asset_ids']
        project_key = value['name'].casefold()
        inspected = sum(1 for asset_id in ids if asset_id in project_inspected_assets[project_key])
        total = len(ids)
        project_rows.append({
            'project': value['name'], 'total': total, 'ever_inspected': inspected,
            'never_inspected': max(total - inspected, 0),
            'coverage_percent': round(inspected * 100 / total) if total else 0,
            'latest_visit_date': latest_by_project.get(value['name'].casefold()),
        })
    project_rows.sort(key=lambda row: (row['coverage_percent'], -row['total'], row['project'].casefold()))

    visit_queryset = _site_visit_aggregate(SiteVisit.objects.all())
    recent_visits = list(visit_queryset.select_related('created_by').order_by('-visit_date', '-created_at')[:8])
    return {
        'smart_asset_count': len(assets),
        'smart_ever_inspected_count': len(ever_ids),
        'smart_never_inspected_count': max(len(assets) - len(ever_ids), 0),
        'smart_repeat_inspected_count': sum(1 for count in successful_counts.values() if count > 1),
        'smart_needs_attention_count': sum(1 for status in latest_result_by_asset.values() if status == SiteInspection.Status.NEEDS_ATTENTION),
        'smart_visit_count': SiteVisit.objects.count(),
        'smart_completed_visit_count': SiteVisit.objects.filter(status=SiteVisit.Status.COMPLETED).count(),
        'smart_open_visit_count': SiteVisit.objects.filter(status=SiteVisit.Status.IN_PROGRESS).count(),
        'smart_project_rows': project_rows[:15],
        'smart_project_count': len(grouped_projects),
        'smart_recent_visits': recent_visits,
    }


def _project_dashboard_context(project, visit_id=None, page_number=None, from_date=None, to_date=None):
    _inspectable = _project_scoped_assets(project).select_related('employee').order_by('asset_type', 'serial_number', 'id')
    project_assets = list(_inspectable)
    asset_ids = [asset.pk for asset in project_assets]
    project_visits = _site_visit_aggregate(SiteVisit.objects.filter(project__iexact=project)).order_by('-visit_date', '-created_at')
    selected_visit = None
    if visit_id:
        selected_visit = project_visits.filter(pk=visit_id).first()
    if selected_visit is None:
        selected_visit = project_visits.first()

    current_status_by_asset = {}
    current_unable_count = 0
    needs_attention_this_visit = 0
    current_additional_count = 0
    if selected_visit:
        current_rows = list(selected_visit.inspections.select_related('asset').all())
        current_status_by_asset = {row.asset_id: row.status for row in current_rows}
        current_unable_count = sum(1 for row in current_rows if row.status == SiteInspection.Status.UNABLE and row.asset_id in asset_ids)
        needs_attention_this_visit = sum(1 for row in current_rows if row.status == SiteInspection.Status.NEEDS_ATTENTION and row.asset_id in asset_ids)
        current_additional_count = sum(1 for row in current_rows if row.is_additional)
    current_inspected_ids = {
        asset_id for asset_id, status in current_status_by_asset.items()
        if asset_id in asset_ids and status in SUCCESS_STATUSES
    }
    total = len(project_assets)
    inspected_this_visit = len(current_inspected_ids)
    never_counts = dict(
        SiteInspection.objects.filter(
            asset_id__in=asset_ids, status__in=SUCCESS_STATUSES,
            visit__project__iexact=project, is_additional=False,
        ).values('asset_id').annotate(total=Count('id')).values_list('asset_id', 'total')
    ) if asset_ids else {}
    history_counts = never_counts
    period_queryset = SiteInspection.objects.filter(
        asset_id__in=asset_ids, status__in=SUCCESS_STATUSES,
        visit__project__iexact=project, is_additional=False,
    )
    if from_date:
        period_queryset = period_queryset.filter(visit__visit_date__gte=from_date)
    if to_date:
        period_queryset = period_queryset.filter(visit__visit_date__lte=to_date)
    period_checked_ids = set(period_queryset.values_list('asset_id', flat=True).distinct())
    history_rows = SiteInspection.objects.filter(
        asset_id__in=asset_ids, status__in=SUCCESS_STATUSES,
        visit__project__iexact=project, is_additional=False,
    ).select_related('visit').order_by('-visit__visit_date', '-inspected_at', '-id')
    last_dates = {}
    for row in history_rows:
        last_dates.setdefault(row.asset_id, row.visit.visit_date)

    table_rows = []
    for asset in project_assets:
        table_rows.append({
            'asset': asset,
            'type_snapshot': _type_snapshot(asset),
            'inspection_count': history_counts.get(asset.pk, 0),
            'last_inspected': last_dates.get(asset.pk),
            'current_status': current_status_by_asset.get(asset.pk, SiteInspection.Status.PENDING),
            'never_inspected': asset.pk not in history_counts,
        })
    pager = Paginator(table_rows, 30)
    assets_page = pager.get_page(page_number)
    return {
        'project': project,
        'project_assets_total': total,
        'project_visit_count': project_visits.count(),
        'project_visits': project_visits,
        'recent_project_visits': project_visits[:12],
        'selected_visit': selected_visit,
        'selected_visit_inspection_total': selected_visit.inspection_total if selected_visit else 0,
        'inspected_this_visit': inspected_this_visit,
        'not_inspected_this_visit': max(total - inspected_this_visit, 0),
        'current_unable_count': current_unable_count,
        'needs_attention_this_visit': needs_attention_this_visit,
        'current_additional_count': current_additional_count,
        'ever_inspected_count': len(history_counts),
        'never_inspected_count': max(total - len(history_counts), 0),
        'repeat_inspected_count': sum(1 for count in history_counts.values() if count > 1),
        'coverage_percent': round(inspected_this_visit * 100 / total) if total else 0,
        'period_checked_count': len(period_checked_ids),
        'not_inspected_in_period': max(total - len(period_checked_ids), 0),
        'period_coverage_percent': round(len(period_checked_ids) * 100 / total) if total else 0,
        'from_date_value': from_date.isoformat() if from_date else '',
        'to_date_value': to_date.isoformat() if to_date else '',
        'asset_rows_page': assets_page,
        'project_names': _project_names(),
    }


def _visit_detail_context(request, visit, posted_row_id=None, posted_row_data=None, form_error=''):
    inspections = visit.inspections.select_related('asset', 'inspected_by').order_by('type_snapshot', 'serial_number_snapshot', 'id')
    q = (request.GET.get('q') or '').strip()
    search_page = None
    if q:
        already_added_ids = visit.inspections.values_list('asset_id', flat=True)
        results = _inspectable_assets().select_related('employee').filter(
            Q(serial_number__icontains=q) | Q(computer_name__icontains=q) |
            Q(computer_name_new__icontains=q) | Q(model__icontains=q) |
            Q(item_type__icontains=q) | Q(employee_name__icontains=q) | Q(hr_code__icontains=q)
        ).order_by('project', 'asset_type', 'serial_number', 'id')
        search_page = Paginator(results, 25).get_page(request.GET.get('search_page'))
        already_ids = set(already_added_ids)
        for asset in search_page.object_list:
            asset.already_in_visit = asset.pk in already_ids
            asset.is_additional_candidate = not _asset_belongs_to_project(asset, visit.project)
    return {
        'visit': visit,
        'inspections': inspections,
        'q': q,
        'search_page': search_page,
        'posted_row_id': posted_row_id,
        'posted_row_data': posted_row_data or {},
        'form_error': form_error,
        'inspection_status_choices': SiteInspection.Status.choices,
        'inspection_count': inspections.count(),
        'inspected_count': inspections.filter(status__in=SUCCESS_STATUSES).count(),
        'pending_count': inspections.filter(status=SiteInspection.Status.PENDING).count(),
        'unable_count': inspections.filter(status=SiteInspection.Status.UNABLE).count(),
    }


@login_required
def _can_manage_visit(user, visit):
    """Only the visit creator (owner) and administrators may cancel/delete it."""
    is_admin = bool(
        getattr(user, 'is_superuser', False)
        or getattr(user, 'is_staff', False)
        or getattr(user, 'is_admin_role', False)
        or getattr(user, 'role', '') == 'ADMIN'
    )
    return is_admin or (visit.created_by_id is not None and visit.created_by_id == user.pk)


def site_visit_list(request):
    query = (request.GET.get('q') or '').strip()
    visits = _site_visit_aggregate(SiteVisit.objects.select_related('created_by'))
    if query:
        visits = visits.filter(Q(project__icontains=query) | Q(technician_name__icontains=query) | Q(notes__icontains=query))
    page = Paginator(visits.order_by('-visit_date', '-created_at'), 30).get_page(request.GET.get('page'))
    return render(request, 'visits/site_visit_list.html', {
        'page': page, 'query': query, 'project_names': _project_names(),
    })


@login_required
def project_inspection_dashboard(request):
    project = (request.GET.get('project') or '').strip()
    if not project:
        return redirect('site_visit_list')
    valid_project = next((item for item in _project_names() if item.casefold() == project.casefold()), None)
    if not valid_project:
        messages.error(request, 'اسم المشروع غير موجود في بيانات SIAC الحالية.')
        return redirect('site_visit_list')
    visit_id = (request.GET.get('visit_id') or '').strip()
    from_text = (request.GET.get('from_date') or '').strip()
    to_text = (request.GET.get('to_date') or '').strip()
    try:
        from_date = date.fromisoformat(from_text) if from_text else None
    except ValueError:
        from_date = None
        messages.error(request, 'تاريخ البداية غير صحيح؛ تم تجاهل الفلتر.')
    try:
        to_date = date.fromisoformat(to_text) if to_text else None
    except ValueError:
        to_date = None
        messages.error(request, 'تاريخ النهاية غير صحيح؛ تم تجاهل الفلتر.')
    if from_date and to_date and from_date > to_date:
        from_date, to_date = to_date, from_date
        messages.warning(request, 'تم عكس تاريخي البداية والنهاية لأن البداية كانت بعد النهاية.')
    context = _project_dashboard_context(
        valid_project, visit_id=visit_id or None, page_number=request.GET.get('page'),
        from_date=from_date, to_date=to_date,
    )
    return render(request, 'visits/project_dashboard.html', context)


@login_required
def site_visit_create(request):
    projects = _project_names()
    initial_project = (request.GET.get('project') or '').strip()
    posted_values = {}
    if request.method == 'POST':
        project = (request.POST.get('project') or '').strip()
        visit_date_text = (request.POST.get('visit_date') or '').strip()
        technician_name = (request.POST.get('technician_name') or '').strip()
        notes = (request.POST.get('notes') or '').strip()
        posted_values = {'project': project, 'visit_date': visit_date_text, 'technician_name': technician_name, 'notes': notes}
        valid_project = next((item for item in projects if item.casefold() == project.casefold()), None)
        try:
            visit_date = date.fromisoformat(visit_date_text)
        except ValueError:
            visit_date = None
        errors = []
        if not valid_project:
            errors.append('اختَر مشروعًا موجودًا في بيانات SIAC.')
        if not visit_date:
            errors.append('تاريخ الزيارة غير صحيح.')
        if not technician_name:
            errors.append('اسم فني IT مطلوب.')
        if len(technician_name) > 255:
            errors.append('اسم الفني أطول من الحد المسموح.')
        if errors:
            for error in errors:
                messages.error(request, error)
        else:
            with transaction.atomic():
                visit = SiteVisit.objects.create(
                    project=valid_project, visit_date=visit_date, technician_name=technician_name,
                    notes=notes, status=SiteVisit.Status.IN_PROGRESS,
                    created_by=request.user,
                )
                seeded = _seed_project_assets(visit)
            log_event(request.user, 'SITE_VISIT_CREATE', request, f'Created site visit id={visit.pk}, project={visit.project}, date={visit.visit_date}, checklist_rows={seeded}')
            messages.success(request, f'تم إنشاء الزيارة وتجهيز قائمة الفحص بعدد {seeded} جهازًا من المشروع.')
            return redirect('site_visit_detail', pk=visit.pk)
        initial_project = project
    return render(request, 'visits/site_visit_form.html', {
        'project_names': projects, 'initial_project': initial_project, 'posted_values': posted_values,
        'today': timezone.localdate().isoformat(),
    })


@login_required
def site_visit_detail(request, pk):
    visit = get_object_or_404(SiteVisit.objects.select_related('created_by'), pk=pk)
    if visit.status == SiteVisit.Status.IN_PROGRESS:
        _seed_project_assets(visit)
        _refresh_peripheral_snapshots(visit)

    if request.method == 'POST':
        action = (request.POST.get('action') or '').strip()
        if action in {'cancel_visit', 'delete_visit'}:
            if not _can_manage_visit(request.user, visit):
                messages.error(request, 'إلغاء أو حذف الزيارة متاح لصاحب الزيارة أو مسؤول النظام فقط.')
                return redirect('site_visit_detail', pk=visit.pk)
            if action == 'cancel_visit':
                if visit.status == SiteVisit.Status.CANCELLED:
                    messages.info(request, 'الزيارة ملغاة بالفعل.')
                else:
                    visit.status = SiteVisit.Status.CANCELLED
                    visit.completed_at = timezone.now()
                    visit.save(update_fields=['status', 'completed_at'])
                    log_event(request.user, 'SITE_VISIT_CANCEL', request, f'Cancelled site visit id={visit.pk}, project={visit.project}, date={visit.visit_date}')
                    messages.success(request, 'تم إلغاء الزيارة. بياناتها محفوظة كسجل ملغى.')
                return redirect('site_visit_detail', pk=visit.pk)
            visit_id, visit_project, visit_date = visit.pk, visit.project, visit.visit_date
            log_event(request.user, 'SITE_VISIT_DELETE', request, f'Deleted site visit id={visit_id}, project={visit_project}, date={visit_date}')
            visit.delete()
            messages.success(request, 'تم حذف الزيارة وقائمة الفحص التابعة لها.')
            return redirect('site_visit_list')

        if visit.status in (SiteVisit.Status.COMPLETED, SiteVisit.Status.CANCELLED):
            messages.error(request, 'الزيارة مكتملة ومقفولة. لا يمكن تغيير سجلها؛ أنشئ زيارة جديدة للتوثيق الجديد.')
            return redirect('site_visit_detail', pk=visit.pk)

        if action == 'add_asset':
            asset_id = (request.POST.get('asset_id') or '').strip()
            asset = get_object_or_404(_inspectable_assets(), pk=asset_id)
            if visit.inspections.filter(asset=asset).exists():
                messages.info(request, 'الجهاز موجود بالفعل في قائمة الزيارة.')
            else:
                try:
                    with transaction.atomic():
                        _create_additional_inspection(visit, asset)
                except IntegrityError:
                    messages.info(request, 'الجهاز أُضيف بالفعل إلى الزيارة من جلسة أخرى.')
                else:
                    log_event(request.user, 'SITE_VISIT_ASSET_ADD', request, f'Added asset id={asset.pk}, serial={asset.serial_number}, to visit id={visit.pk}; source project={asset.project or "(blank)"}')
                    messages.success(request, 'تمت إضافة الجهاز إلى الزيارة من غير تعديل بياناته الأساسية في SIAC.')
            return redirect('site_visit_detail', pk=visit.pk)

        if action == 'complete_visit':
            _seed_project_assets(visit)
            pending_count = visit.inspections.filter(status=SiteInspection.Status.PENDING).count()
            if pending_count:
                messages.error(request, f'لا يمكن اعتماد الزيارة: ما زال هناك {pending_count} جهازًا لم تُسجّل له نتيجة فحص أو تعذّر فحص.')
            else:
                visit.status = SiteVisit.Status.COMPLETED
                visit.completed_at = timezone.now()
                visit.save(update_fields=['status', 'completed_at'])
                log_event(request.user, 'SITE_VISIT_COMPLETE', request, f'Completed site visit id={visit.pk}, project={visit.project}, rows={visit.inspections.count()}')
                messages.success(request, 'تم اعتماد الزيارة وإغلاقها بنجاح.')
            return redirect('site_visit_detail', pk=visit.pk)

        if action == 'save_inspection':
            inspection = get_object_or_404(SiteInspection.objects.select_related('asset'), pk=request.POST.get('inspection_id'), visit=visit)
            values = {field: (request.POST.get(field) or '').strip() for field in EDITABLE_SNAPSHOT_FIELDS}
            reasons = {field: (request.POST.get(f'reason_{field}') or '').strip() for field in EDITABLE_SNAPSHOT_FIELDS}
            comments = (request.POST.get('comments') or '').strip()
            new_status = (request.POST.get('status') or SiteInspection.Status.PENDING).strip()
            error_list = []
            changed = {}
            for field, (label, reason_label) in EDITABLE_SNAPSHOT_FIELDS.items():
                if len(values[field]) > 255:
                    error_list.append(f'{label}: الحد الأقصى 255 حرفًا.')
                if values[field] != (getattr(inspection, field) or ''):
                    changed[field] = values[field]
                    if not reasons[field]:
                        error_list.append(f'اكتب سبب التعديل في {label}.')
            if new_status not in dict(SiteInspection.Status.choices):
                error_list.append('حالة الفحص غير صحيحة.')
            if inspection.status != SiteInspection.Status.PENDING and new_status == SiteInspection.Status.PENDING:
                error_list.append('لا يمكن إعادة فحص مكتمل إلى «لم يتم الفحص». سجّل النتيجة الصحيحة أو تعذر الفحص مع توضيح السبب.')
            if new_status == SiteInspection.Status.UNABLE and not comments:
                error_list.append('اكتب سبب تعذر الفحص في COMMENTS.')
            if new_status == SiteInspection.Status.NEEDS_ATTENTION and not comments:
                error_list.append('اكتب الملاحظة أو سبب المتابعة في COMMENTS لهذه الحالة.')
            if new_status == SiteInspection.Status.PENDING:
                if comments or changed:
                    error_list.append('حدد نتيجة الفحص أو تعذر الفحص قبل حفظ الملاحظات والتعديلات.')
                else:
                    error_list.append('اختَر نتيجة الفحص قبل الحفظ: سليم، يحتاج متابعة، أو تعذر الفحص.')
            if len(comments) > 10000:
                error_list.append('الملاحظات أطول من الحد المسموح (10000 حرف).')
            if error_list:
                context = _visit_detail_context(request, visit, posted_row_id=inspection.pk,
                    posted_row_data={**values, 'status': new_status, 'comments': comments, 'reasons': reasons},
                    form_error=' '.join(error_list))
                return render(request, 'visits/site_visit_detail.html', context, status=400)

            auto_entries = []
            effective_day = timezone.localdate().isoformat()
            for field, new_value in changed.items():
                label, _reason_label = EDITABLE_SNAPSHOT_FIELDS[field]
                old_value = getattr(inspection, field) or ''
                reason = reasons[field]
                auto_entries.append(
                    f'[AUTO {effective_day}] {label} changed from "{old_value or "(blank)"}" to "{new_value or "(blank)"}". Reason: {reason}'
                )
            with transaction.atomic():
                locked = SiteInspection.objects.select_for_update().get(pk=inspection.pk)
                # Validate edits against latest saved values to avoid overwriting a concurrent technician update.
                conflict_fields = [
                    field for field in EDITABLE_SNAPSHOT_FIELDS
                    if getattr(locked, field) != getattr(inspection, field)
                ]
                if (locked.status != inspection.status or locked.comments != inspection.comments or
                        locked.auto_comments != inspection.auto_comments or
                        locked.inspected_at != inspection.inspected_at or
                        locked.inspected_by_id != inspection.inspected_by_id):
                    conflict_fields.append('status/comments')
                if conflict_fields:
                    messages.error(request, 'تم تعديل السجل من مستخدم آخر. راجع البيانات الحالية ثم أعد المحاولة.')
                    return redirect('site_visit_detail', pk=visit.pk)
                for field, value in values.items():
                    setattr(locked, field, value)
                locked.comments = comments
                locked.status = new_status
                if auto_entries:
                    locked.auto_comments = '\n'.join(part for part in [locked.auto_comments.strip(), *auto_entries] if part)
                if new_status == SiteInspection.Status.PENDING:
                    locked.inspected_by = None
                    locked.inspected_at = None
                else:
                    locked.inspected_by = request.user
                    locked.inspected_at = timezone.now()
                locked.save()
            change_desc = f'Updated inspection id={inspection.pk}, visit id={visit.pk}, asset id={inspection.asset_id}, status={new_status}, changed_fields={",".join(changed.keys()) or "none"}'
            log_event(request.user, 'SITE_INSPECTION_UPDATE', request, change_desc)
            messages.success(request, 'تم حفظ نتيجة الفحص. كل تعديل وسببه محفوظان داخل COMMENTS، ولم يتم تعديل بيانات SIAC الأساسية.')
            return redirect('site_visit_detail', pk=visit.pk)

        messages.error(request, 'الإجراء المطلوب غير معروف.')
        return redirect('site_visit_detail', pk=visit.pk)

    context = _visit_detail_context(request, visit)
    context['can_manage_visit'] = _can_manage_visit(request.user, visit)
    return render(request, 'visits/site_visit_detail.html', context)


@login_required
def site_visit_export(request, pk):
    visit = get_object_or_404(SiteVisit, pk=pk)
    inspections = visit.inspections.select_related('asset').order_by('type_snapshot', 'serial_number_snapshot', 'id')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Periodic Inspection'
    sheet.append(EXCEL_HEADERS)
    for inspection in inspections:
        sheet.append([
            visit.project,
            visit.visit_date.isoformat(),
            visit.technician_name,
            inspection.user_name,
            inspection.type_snapshot,
            inspection.model_snapshot,
            inspection.serial_number_snapshot,
            inspection.computer_name_snapshot,
            inspection.monitor_model,
            inspection.monitor_serial,
            inspection.printer_model,
            inspection.printer_serial,
            inspection.combined_comments,
        ])
    header_fill = PatternFill('solid', fgColor='17365D')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    widths = [26, 14, 24, 26, 16, 26, 24, 26, 26, 26, 26, 26, 55]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    log_event(request.user, 'SITE_VISIT_EXPORT', request, f'Exported inspection workbook for visit id={visit.pk}, project={visit.project}, row_count={inspections.count()}')
    filename_project = ''.join(char if char.isalnum() or char in '-_' else '_' for char in visit.project)[:50] or 'Project'
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="SIAC_Site_Inspection_{filename_project}_{visit.visit_date.isoformat()}.xlsx"'
    return response
