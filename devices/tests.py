from datetime import date

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from django.urls import reverse

from audit.models import AuditLog
from .models import Asset, AssetMovement, Employee, SiteInspection, SiteVisit


User = get_user_model()


class CustodyMovementWorkflowTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username='siacadmin', password='Secure-test-password-123', role='ADMIN'
        )
        self.regular = User.objects.create_user(
            username='siacuser', password='Secure-test-password-123', role='USER'
        )
        self.employee1 = Employee.objects.create(hr_code='1001', name='Employee One')
        self.employee2 = Employee.objects.create(hr_code='1002', name='Employee Two')
        self.asset = Asset.objects.create(
            asset_type=Asset.Types.COMPUTER,
            asset_status=Asset.Status.IN_STOCK,
            serial_number='SN-TEST-100',
            brand='Dell', model='OptiPlex 7010', item_type='Desktop PC',
            source_sheet='manual',
        )
        self.client.force_login(self.admin)

    def test_assign_then_transfer_then_return_preserves_all_events(self):
        response = self.client.post(reverse('movement_create', args=[self.asset.pk]), {
            'event_type': AssetMovement.EventTypes.ASSIGNED,
            'target_hr_code': self.employee1.hr_code,
            'effective_date': '2026-10-17',
            'reason': 'تسليم جهاز للموظف',
            'reference_number': 'HANDOVER-1',
            'notes': 'اختبار حركة',
        })
        self.assertEqual(response.status_code, 302)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.hr_code, self.employee1.hr_code)
        self.assertEqual(self.asset.asset_status, Asset.Status.ACTIVE)
        self.assertEqual(AssetMovement.objects.count(), 1)
        self.assertEqual(AssetMovement.objects.get().created_by, self.admin)
        self.assertTrue(AuditLog.objects.filter(action='ASSET_ASSIGN').exists())

        response = self.client.post(reverse('movement_create', args=[self.asset.pk]), {
            'event_type': AssetMovement.EventTypes.TRANSFERRED,
            'target_hr_code': self.employee2.hr_code,
            'effective_date': '2026-10-18',
            'reason': 'نقل عهدة', 'reference_number': 'TRANSFER-1', 'notes': '',
        })
        self.assertEqual(response.status_code, 302)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.hr_code, self.employee2.hr_code)
        self.assertEqual(AssetMovement.objects.count(), 2)
        latest_transfer = AssetMovement.objects.order_by('-id').first()
        self.assertEqual(latest_transfer.from_hr_code, self.employee1.hr_code)
        self.assertEqual(latest_transfer.to_hr_code, self.employee2.hr_code)
        self.assertTrue(AuditLog.objects.filter(action='ASSET_TRANSFER').exists())

        response = self.client.post(reverse('movement_create', args=[self.asset.pk]), {
            'event_type': AssetMovement.EventTypes.RETURNED,
            'target_hr_code': '', 'effective_date': '2026-10-19',
            'reason': 'إرجاع للمخزن', 'reference_number': 'RETURN-1', 'notes': '',
        })
        self.assertEqual(response.status_code, 302)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.hr_code, '')
        self.assertEqual(self.asset.asset_status, Asset.Status.IN_STOCK)
        self.assertEqual(AssetMovement.objects.count(), 3)
        self.assertTrue(AuditLog.objects.filter(action='ASSET_RETURN').exists())

    def test_scrap_requires_number_and_records_holder_then_scrap_state(self):
        self.asset.hr_code = self.employee1.hr_code
        self.asset.employee = self.employee1
        self.asset.employee_name = self.employee1.name
        self.asset.asset_status = Asset.Status.ACTIVE
        self.asset.save()

        response = self.client.post(reverse('movement_create', args=[self.asset.pk]), {
            'event_type': AssetMovement.EventTypes.SCRAPPED,
            'effective_date': '2026-10-17', 'reason': 'تخريد رسمي',
            'reference_number': 'SCRAP-77', 'notes': 'محضر تخريد معتمد',
        })
        self.assertEqual(response.status_code, 302)
        self.asset.refresh_from_db()
        self.assertEqual(self.asset.asset_status, Asset.Status.SCRAPPED)
        self.assertEqual(self.asset.scrap_number, 'SCRAP-77')
        self.assertEqual(self.asset.scrap_date, '2026-10-17')
        movement = AssetMovement.objects.get(event_type=AssetMovement.EventTypes.SCRAPPED)
        self.assertEqual(movement.from_hr_code, self.employee1.hr_code)
        self.assertEqual(movement.to_employee_name, 'التخريد')
        self.assertTrue(AuditLog.objects.filter(action='ASSET_SCRAP').exists())

    def test_regular_user_cannot_open_admin_movement_or_dashboard_pages(self):
        self.client.force_login(self.regular)
        for url in [reverse('movement_list'), reverse('movement_asset_search'), reverse('inventory_dashboard'), reverse('inventory_dashboard_export')]:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 302)

    def test_dashboard_filters_model_counts(self):
        Asset.objects.create(
            asset_type=Asset.Types.COMPUTER, asset_status=Asset.Status.ACTIVE,
            serial_number='SN-LAPTOP-1', brand='HP', model='EliteBook 840', item_type='Laptop',
            hr_code=self.employee1.hr_code, employee=self.employee1, employee_name=self.employee1.name,
            source_sheet='manual',
        )
        response = self.client.get(reverse('inventory_dashboard'), {'model_type': 'laptop', 'model_q': 'EliteBook'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'EliteBook 840')
        self.assertContains(response, 'Laptop')


class SiteVisitInspectionWorkflowTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username='visitadmin', password='Secure-test-password-123', role='ADMIN'
        )
        self.regular = User.objects.create_user(
            username='visittech', password='Secure-test-password-123', role='USER'
        )
        self.computer = Asset.objects.create(
            asset_type=Asset.Types.COMPUTER, asset_status=Asset.Status.ACTIVE,
            source_sheet='manual', project='Project Alpha', serial_number='PC-ALPHA-01',
            item_type='Desktop PC', model='OptiPlex 7010', computer_name='ALPHA-PC-01',
            employee_name='Original User', hr_code='',
        )
        self.monitor = Asset.objects.create(
            asset_type=Asset.Types.MONITOR, asset_status=Asset.Status.ACTIVE,
            source_sheet='manual', project='Project Alpha', serial_number='MON-ALPHA-01',
            model='Dell P2422H', employee_name='Original User', hr_code='',
        )
        self.printer = Asset.objects.create(
            asset_type=Asset.Types.PRINTER, asset_status=Asset.Status.IN_STOCK,
            source_sheet='manual', project='Project Alpha', serial_number='PRN-ALPHA-01',
            model='HP LaserJet', employee_name='', hr_code='',
        )
        self.other_project_asset = Asset.objects.create(
            asset_type=Asset.Types.COMPUTER, asset_status=Asset.Status.ACTIVE,
            source_sheet='manual', project='Project Beta', serial_number='PC-BETA-01',
            item_type='Laptop', model='EliteBook 840', computer_name='BETA-LAPTOP-01',
        )
        self.client.force_login(self.regular)

    def test_create_visit_builds_checklist_for_all_project_hardware(self):
        response = self.client.post(reverse('site_visit_create'), {
            'project': 'Project Alpha', 'visit_date': '2026-10-18',
            'technician_name': 'IT Tech One', 'notes': 'Quarterly check',
        })
        self.assertEqual(response.status_code, 302)
        visit = SiteVisit.objects.get(project='Project Alpha')
        self.assertEqual(visit.inspections.count(), 3)
        self.assertSetEqual(
            set(visit.inspections.values_list('asset_id', flat=True)),
            {self.computer.pk, self.monitor.pk, self.printer.pk},
        )
        self.assertEqual(visit.status, SiteVisit.Status.IN_PROGRESS)

    def test_visit_creator_can_cancel_and_delete_visit(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.regular,
        )
        url = reverse('site_visit_detail', args=[visit.pk])
        response = self.client.post(url, {'action': 'cancel_visit'})
        self.assertEqual(response.status_code, 302)
        visit.refresh_from_db()
        self.assertEqual(visit.status, SiteVisit.Status.CANCELLED)

        response = self.client.post(url, {'action': 'delete_visit'})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(SiteVisit.objects.filter(pk=visit.pk).exists())

    def test_non_creator_non_admin_cannot_delete_visit(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.admin,
        )
        response = self.client.post(reverse('site_visit_detail', args=[visit.pk]), {
            'action': 'delete_visit',
        })
        self.assertEqual(response.status_code, 403)
        self.assertTrue(SiteVisit.objects.filter(pk=visit.pk).exists())

    def test_non_creator_non_admin_cannot_cancel_visit(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.admin,
        )
        response = self.client.post(reverse('site_visit_detail', args=[visit.pk]), {
            'action': 'cancel_visit',
        })
        self.assertEqual(response.status_code, 403)
        visit.refresh_from_db()
        self.assertNotEqual(visit.status, SiteVisit.Status.CANCELLED)

    def test_global_search_adds_other_project_asset_without_changing_asset(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.regular,
        )
        response = self.client.post(reverse('site_visit_detail', args=[visit.pk]), {
            'action': 'add_asset', 'asset_id': self.other_project_asset.pk,
        })
        self.assertEqual(response.status_code, 302)
        inspection = SiteInspection.objects.get(visit=visit, asset=self.other_project_asset)
        self.assertTrue(inspection.is_additional)
        self.other_project_asset.refresh_from_db()
        self.assertEqual(self.other_project_asset.project, 'Project Beta')

    def test_snapshot_correction_requires_reason_and_never_updates_asset_table(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.regular,
        )
        from .site_visits import _seed_project_assets
        _seed_project_assets(visit)
        inspection = SiteInspection.objects.get(visit=visit, asset=self.computer)
        url = reverse('site_visit_detail', args=[visit.pk])
        base_data = {
            'action': 'save_inspection', 'inspection_id': inspection.pk,
            'user_name': 'Observed User', 'monitor_model': inspection.monitor_model,
            'monitor_serial': inspection.monitor_serial, 'printer_model': inspection.printer_model,
            'printer_serial': inspection.printer_serial, 'status': SiteInspection.Status.INSPECTED,
            'comments': 'Checked at site',
        }
        response = self.client.post(url, base_data)
        self.assertEqual(response.status_code, 400)
        inspection.refresh_from_db()
        self.assertEqual(inspection.user_name, 'Original User')

        base_data['reason_user_name'] = 'Employee name verified on site'
        for field in ('monitor_model', 'monitor_serial', 'printer_model', 'printer_serial'):
            base_data[f'reason_{field}'] = ''
        response = self.client.post(url, base_data)
        self.assertEqual(response.status_code, 302)
        inspection.refresh_from_db()
        self.assertEqual(inspection.user_name, 'Observed User')
        self.assertIn('Employee name verified on site', inspection.auto_comments)
        self.assertEqual(inspection.status, SiteInspection.Status.INSPECTED)
        self.computer.refresh_from_db()
        self.assertEqual(self.computer.employee_name, 'Original User')

    def test_additional_inspection_does_not_inflate_source_project_coverage(self):
        from .site_visits import _project_dashboard_context, site_visit_global_metrics
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.regular,
        )
        SiteInspection.objects.create(
            visit=visit, asset=self.other_project_asset, is_additional=True,
            source_project='Project Beta', serial_number_snapshot='PC-BETA-01',
            type_snapshot='LAPTOP', model_snapshot='EliteBook 840',
            status=SiteInspection.Status.INSPECTED, inspected_by=self.regular,
            inspected_at=timezone.now(),
        )
        beta_context = _project_dashboard_context('Project Beta')
        self.assertEqual(beta_context['ever_inspected_count'], 0)
        self.assertEqual(beta_context['coverage_percent'], 0)
        beta_row = next(row for row in site_visit_global_metrics()['smart_project_rows'] if row['project'] == 'Project Beta')
        self.assertEqual(beta_row['ever_inspected'], 0)
        self.assertEqual(beta_row['coverage_percent'], 0)

    def test_project_dashboard_and_export_work_for_regular_user(self):
        visit = SiteVisit.objects.create(
            project='Project Alpha', visit_date=date(2026, 10, 18),
            technician_name='IT Tech One', created_by=self.regular,
        )
        from .site_visits import _seed_project_assets
        _seed_project_assets(visit)
        inspection = visit.inspections.get(asset=self.computer)
        inspection.status = SiteInspection.Status.INSPECTED
        inspection.inspected_at = timezone.now()
        inspection.save()
        response = self.client.get(reverse('project_inspection_dashboard'), {'project': 'Project Alpha'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Smart IT Asset Dashboard')
        export_response = self.client.get(reverse('site_visit_export', args=[visit.pk]))
        self.assertEqual(export_response.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(export_response.content), read_only=True)
        self.assertEqual(
            list(next(workbook.active.iter_rows(values_only=True))),
            ['Project', 'visit date', 'IT TECH NAME', 'USER NAME', 'TYPE', 'MODEL',
             'SERIAL NUMBER', 'COMPUTER NAME', 'Monitor Model', 'Monitor Serial',
             'printer model', 'printer serial', 'COMMENTS'],
        )
