import io
import json
from collections import OrderedDict
from datetime import datetime

from django.http import HttpResponse
from django.db.models import F
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .models import Asset, AssetMovement, Employee

NAVY = '17365D'
TEAL = '1F8A8A'
LIGHT = 'EAF1F8'
WHITE = 'FFFFFF'

EMPLOYEE_COLUMNS = [
    ('HR Code', 'hr_code'), ('Employee Name', 'name'), ('Arabic Name', 'arabic_name'),
    ('Company', 'company'), ('Cost Center', 'cost_center'), ('Project / Site', 'project'), ('Sector', 'sector'),
    ('Department', 'department'), ('Section', 'section'), ('Unit', 'unit'),
    ('Office', 'office'), ('Team', 'team'), ('Position', 'position'),
    ('Employment Status', 'employment_status'), ('Hiring Date', 'hiring_date'), ('Termination Date', 'termination_date'),
]
ASSET_COLUMNS = [
    ('HR Code', 'hr_code'), ('Employee Name', 'employee_name'), ('Asset Type', 'get_asset_type_display'),
    ('Asset Status', 'get_asset_status_display'), ('Company', 'company'), ('Project', 'project'),
    ('Sector', 'sector'), ('Department', 'department'), ('Computer Name New', 'computer_name_new'),
    ('Computer Name', 'computer_name'), ('Item / Type', 'item_type'), ('Brand', 'brand'),
    ('Model', 'model'), ('Serial Number', 'serial_number'), ('Asset Number', 'asset_tag'),
    ('Manufacturing Year', 'manufacturing_year'), ('Processor', 'processor'), ('RAM', 'ram'),
    ('Storage / HDD', 'storage'), ('Email', 'email'), ('Toner Code', 'toner_code'),
    ('License Version', 'license_version'), ('License Activated', 'license_activated'),
    ('Purchase Date', 'purchase_date'), ('Scrap Date', 'scrap_date'), ('Scrap Number', 'scrap_number'),
    ('Notes', 'notes'), ('Source Sheet', 'source_sheet'), ('Source Row', 'source_row'),
    ('Last Updated', 'updated_at'),
]

RAW_OVERRIDES = {
    'HR Code': 'hr_code', 'Employee Code': 'hr_code', 'HR ID': 'hr_code',
    'HR ID (clean)': 'hr_code', 'Resolved HR ID': 'hr_code',
    'English Nmae': 'employee_name', 'Employee Name': 'employee_name', 'Name': 'employee_name',
    'User Name': 'employee_name', 'Username': 'employee_name', 'Original Name': 'employee_name',
    'Company': 'company', 'Project': 'project', 'Original Project': 'project', 'Location': 'project',
    'Sector': 'sector', 'Original Sector': 'sector', 'Sector (from ID)': 'sector',
    'Department': 'department', 'Original Department': 'department',
    'Computer Name New': 'computer_name_new', 'Computer Name': 'computer_name',
    'Serial Number': 'serial_number', 'Serial': 'serial_number', 'Computer SN': 'serial_number',
    'Item': 'item_type', 'Type': 'item_type', 'Machin Type': 'item_type', 'Gstar PRO': 'item_type',
    'Brand': 'brand', 'Model': 'model', 'manufacturing Year': 'manufacturing_year',
    'manufacturing year': 'manufacturing_year', 'processor': 'processor', 'ram': 'ram', 'hdd': 'storage',
    'Email': 'email', 'Email to send License': 'email', 'Toner Code': 'toner_code',
    'Version': 'license_version', 'Activated': 'license_activated', 'Asset Number': 'asset_tag',
    'Date': 'purchase_date', 'تاريخ التخريد': 'scrap_date', 'رقم التخريد': 'scrap_number',
}


def _value(obj, attr):
    val = getattr(obj, attr, '')
    return val() if callable(val) else val


def _safe_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if hasattr(value, 'isoformat'):
        try:
            return value.isoformat(sep=' ', timespec='seconds') if isinstance(value, datetime) else value.isoformat()
        except TypeError:
            return value.isoformat()
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        # Prevent text from being interpreted as an Excel formula when exported.
        return "'" + value
    return value


def _format_sheet(ws):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=Side(style='medium', color=TEAL))
    ws.row_dimensions[1].height = 32
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in list(col_cells)[:200]:
            text = '' if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(text), 50))
            if cell.row > 1:
                cell.alignment = Alignment(vertical='top', wrap_text=False)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 34)


def _raw_value(asset, header):
    key = RAW_OVERRIDES.get(header)
    if key:
        val = getattr(asset, key, None)
        if key == 'employee_name' and not val and asset.employee_id:
            val = asset.employee.name
        # If the admin intentionally cleared a mapped field, export the cleared value
        # rather than resurrecting an older value from the imported source JSON.
        return '' if val is None else val
    return asset.raw_data.get(header, '')


def build_inventory_excel_response():
    wb = Workbook()
    ws_emp = wb.active
    ws_emp.title = 'Employees'
    ws_emp.append([label for label, _ in EMPLOYEE_COLUMNS])
    for employee in Employee.objects.order_by('hr_code').iterator(chunk_size=500):
        ws_emp.append([_safe_excel_value(_value(employee, attr)) for _, attr in EMPLOYEE_COLUMNS])
    _format_sheet(ws_emp)

    ws_assets = wb.create_sheet('Asset Register')
    ws_assets.append([label for label, _ in ASSET_COLUMNS])
    assets = list(Asset.objects.select_related('employee').order_by('source_sheet', 'source_row', 'id'))
    for asset in assets:
        ws_assets.append([_safe_excel_value(_value(asset, attr)) for _, attr in ASSET_COLUMNS])
    _format_sheet(ws_assets)

    grouped = OrderedDict()
    for asset in assets:
        grouped.setdefault(asset.source_sheet or 'manual', []).append(asset)
    ws_moves = wb.create_sheet('Movement History')
    movement_headers = [
        'Event ID', 'Serial Number', 'Asset Type', 'Event Type', 'Effective Date',
        'From HR Code', 'From Employee', 'To HR Code', 'To Employee',
        'From Status', 'To Status', 'Reason', 'Reference Number', 'Notes',
        'Source Sheet', 'Source Row', 'Recorded At', 'Recorded By',
    ]
    ws_moves.append(movement_headers)
    movements = AssetMovement.objects.select_related('asset', 'from_employee', 'to_employee', 'created_by').order_by(F('effective_date').asc(nulls_first=True), 'created_at', 'id')
    for movement in movements.iterator(chunk_size=500):
        ws_moves.append([
            movement.pk,
            movement.asset.serial_number,
            movement.asset.get_asset_type_display(),
            movement.get_event_type_display(),
            _safe_excel_value(movement.effective_date),
            movement.from_hr_code,
            movement.from_employee_name or (movement.from_employee.name if movement.from_employee_id else ''),
            movement.to_hr_code,
            movement.to_employee_name or (movement.to_employee.name if movement.to_employee_id else ''),
            movement.from_status,
            movement.to_status,
            movement.reason,
            movement.reference_number,
            movement.notes,
            movement.source_sheet,
            movement.source_row,
            _safe_excel_value(movement.created_at),
            movement.created_by.get_username() if movement.created_by_id else '',
        ])
    _format_sheet(ws_moves)

    for original_sheet, sheet_assets in grouped.items():
        # Create a tab per original source sheet, retaining its source-specific columns.
        base_name = {'all': 'Computers', 'manual': 'Manual Entry'}.get(original_sheet, original_sheet)
        base_name = base_name[:31]
        sheet_name = base_name
        suffix = 2
        while sheet_name in wb.sheetnames:
            tail = f' {suffix}'
            sheet_name = base_name[:31-len(tail)] + tail
            suffix += 1
        ws = wb.create_sheet(sheet_name)
        raw_headers = []
        seen = set()
        for asset in sheet_assets:
            for header in (asset.raw_data or {}).keys():
                normalized = str(header).strip()
                key = normalized.casefold()
                if normalized and key not in seen:
                    seen.add(key)
                    raw_headers.append(normalized)
        headers = list(raw_headers)
        for label, _attr in ASSET_COLUMNS:
            if label.casefold() not in seen:
                headers.append(label)
                seen.add(label.casefold())
        ws.append(headers)
        for asset in sheet_assets:
            ws.append([_safe_excel_value(_raw_value(asset, header)) for header in headers])
        _format_sheet(ws)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="SIAC_Inventory_Export_{datetime.now():%Y%m%d_%H%M}.xlsx"'
    return response
