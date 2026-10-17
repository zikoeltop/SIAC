from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q, Count, F, Prefetch, Case, When, Value, CharField
from django.http import FileResponse, HttpResponse
from django.core.paginator import Paginator
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.text import slugify
from django.utils import timezone
from openpyxl import Workbook
import io

from accounts.decorators import admin_required
from audit.models import AuditLog
from audit.utils import log_event
from .forms import (AssetForm, EmployeeForm, ExcelUploadForm, InventoryWorkbookImportForm, CustodyMovementForm)
from .inventory_exports import build_inventory_excel_response
from .inventory_services import import_inventory_workbook, preview_inventory_workbook
from .models import Asset, AssetMovement, Device, DeviceImport, Employee, PendingExcelImport, PendingInventoryImport
from .services import create_pending_import, import_excel, normalize_serial
from .pdf import build_pdf, build_custody_pdf


def _custody_assets(hr_code):
    if not hr_code:
        return Asset.objects.none()
    return Asset.objects.filter(hr_code__iexact=hr_code, asset_status=Asset.Status.ACTIVE).exclude(
        asset_type=Asset.Types.SCRAP
    ).select_related('employee').order_by('asset_type', 'serial_number', 'id')


def _asset_person_snapshot(asset):
    employee = getattr(asset, 'employee', None)
    return {
        'employee': employee if getattr(asset, 'employee_id', None) else None,
        'hr_code': (getattr(asset, 'hr_code', '') or '').strip(),
        'name': (employee.name if employee else getattr(asset, 'employee_name', '') or '').strip(),
        'status': getattr(asset, 'asset_status', '') or '',
    }


def _record_asset_movement(request, after, form, before=None):
    """Record custody/state changes as append-only history; never overwrite prior events."""
    old = _asset_person_snapshot(before) if before is not None else {
        'employee': None, 'hr_code': '', 'name': 'المخزون', 'status': ''
    }
    new = _asset_person_snapshot(after)
    old_hr = old['hr_code'].casefold()
    new_hr = new['hr_code'].casefold()
    old_status, new_status = old['status'], new['status']

    if new_status == Asset.Status.SCRAPPED and old_status != Asset.Status.SCRAPPED:
        event_type = AssetMovement.EventTypes.SCRAPPED
        to_employee, to_hr, to_name = None, '', 'التخريد'
        reason = (form.cleaned_data.get('movement_reason') or '').strip() or 'تخريد الأصل'
    elif old_hr != new_hr:
        if old_hr and new_hr:
            event_type = AssetMovement.EventTypes.TRANSFERRED
        elif new_hr:
            event_type = AssetMovement.EventTypes.ASSIGNED
        else:
            event_type = AssetMovement.EventTypes.RETURNED
        to_employee, to_hr, to_name = new['employee'], new['hr_code'], new['name'] or new['hr_code'] or 'المخزون'
        reason = (form.cleaned_data.get('movement_reason') or '').strip() or event_type.label
    elif before is not None and old_status != new_status:
        event_type = AssetMovement.EventTypes.STATUS_CHANGED
        to_employee, to_hr, to_name = new['employee'], new['hr_code'], new['name'] or new['hr_code'] or 'المخزون'
        reason = (form.cleaned_data.get('movement_reason') or '').strip() or 'تغيير حالة الأصل'
    elif before is None:
        if new_status == Asset.Status.SCRAPPED:
            event_type = AssetMovement.EventTypes.SCRAPPED
            to_employee, to_hr, to_name = None, '', 'التخريد'
            reason = (form.cleaned_data.get('movement_reason') or '').strip() or 'إضافة سجل جهاز مخرد'
        elif new_hr:
            event_type = AssetMovement.EventTypes.ASSIGNED
            to_employee, to_hr, to_name = new['employee'], new['hr_code'], new['name'] or new['hr_code']
            reason = (form.cleaned_data.get('movement_reason') or '').strip() or 'تسجيل عهدة جديدة'
        else:
            event_type = AssetMovement.EventTypes.ADDED_TO_STOCK
            to_employee, to_hr, to_name = None, '', 'المخزون'
            reason = (form.cleaned_data.get('movement_reason') or '').strip() or 'إضافة أصل إلى المخزون'
    else:
        return None

    movement_date = form.cleaned_data.get('movement_date') or timezone.localdate()
    movement = AssetMovement.objects.create(
        asset=after,
        event_type=event_type,
        effective_date=movement_date,
        from_employee=old['employee'],
        to_employee=to_employee,
        from_hr_code=old['hr_code'],
        from_employee_name=old['name'],
        to_hr_code=to_hr,
        to_employee_name=to_name,
        from_status=old_status,
        to_status=new_status,
        reason=reason,
        reference_number=(after.scrap_number or '') if event_type == AssetMovement.EventTypes.SCRAPPED else '',
        notes='تم تسجيل الحركة من شاشة إدارة الأصول بواسطة المستخدم المسؤول.' if request.user.is_authenticated else '',
        created_by=request.user if request.user.is_authenticated else None,
    )
    return movement


def _movement_history_groups(assets):
    """Group movement rows without triggering one extra SQL query per asset."""
    groups = []
    for asset in assets:
        cached = getattr(asset, '_prefetched_objects_cache', {}).get('movements')
        if cached is not None:
            movements = list(cached)
        else:
            movements = list(asset.movements.select_related('from_employee', 'to_employee', 'created_by').order_by(F('effective_date').asc(nulls_first=True), 'created_at', 'id'))
        if movements:
            groups.append({'asset': asset, 'movements': movements})
    return groups


def _with_movement_history(queryset):
    ordered_movements = AssetMovement.objects.select_related(
        'from_employee', 'to_employee', 'created_by'
    ).order_by(F('effective_date').asc(nulls_first=True), 'created_at', 'id')
    return queryset.prefetch_related(Prefetch('movements', queryset=ordered_movements))


def _omit_duplicate_scrap_rows(assets):
    """If a Scraped source row matches a physical asset serial, show history under the physical asset."""
    assets = list(assets)
    hardware_types = {Asset.Types.COMPUTER, Asset.Types.PRINTER, Asset.Types.MONITOR, Asset.Types.OTHER}
    physical_serials = {a.serial_number.casefold() for a in assets if a.serial_number and a.asset_type in hardware_types}
    return [a for a in assets if not (
        a.asset_type == Asset.Types.SCRAP and a.serial_number and a.serial_number.casefold() in physical_serials
    )]


def _tracked_assets():
    """Canonical asset rows, excluding workbook-only pipeline and duplicate Scraped rows."""
    return Asset.objects.exclude(asset_type__in=[
        Asset.Types.SCRAP, Asset.Types.NEW_2026, Asset.Types.REQUEST_2025,
    ])


def _hardware_assets():
    return Asset.objects.filter(asset_type__in=[
        Asset.Types.COMPUTER, Asset.Types.MONITOR, Asset.Types.PRINTER, Asset.Types.OTHER,
    ])


def _dashboard_model_rows(request):
    """Filtered PC/laptop/monitor/printer counts, grouped by brand and model."""
    model_type = (request.GET.get('model_type') or 'all').strip()
    if model_type not in {'all', 'pc', 'laptop', 'monitor', 'printer'}:
        model_type = 'all'
    status_filter = (request.GET.get('status') or 'all').strip()
    if status_filter not in {'all', 'active', 'in_stock', 'pending', 'scrapped', 'archived'}:
        status_filter = 'all'
    model_query = (request.GET.get('model_q') or '').strip()

    qs = Asset.objects.filter(asset_type__in=[Asset.Types.COMPUTER, Asset.Types.MONITOR, Asset.Types.PRINTER])
    laptop_q = (
        Q(item_type__icontains='laptop') | Q(model__icontains='laptop') |
        Q(item_type__icontains='notebook') | Q(model__icontains='notebook')
    )
    if model_type == 'pc':
        qs = qs.filter(asset_type=Asset.Types.COMPUTER).exclude(laptop_q)
    elif model_type == 'laptop':
        qs = qs.filter(asset_type=Asset.Types.COMPUTER).filter(laptop_q)
    elif model_type == 'monitor':
        qs = qs.filter(asset_type=Asset.Types.MONITOR)
    elif model_type == 'printer':
        qs = qs.filter(asset_type=Asset.Types.PRINTER)

    if status_filter == 'active':
        qs = qs.filter(asset_status=Asset.Status.ACTIVE).exclude(hr_code='')
    elif status_filter == 'in_stock':
        qs = qs.filter(Q(asset_status=Asset.Status.IN_STOCK) | Q(asset_status=Asset.Status.ACTIVE, hr_code=''))
    elif status_filter == 'pending':
        qs = qs.filter(asset_status=Asset.Status.PENDING)
    elif status_filter == 'scrapped':
        qs = qs.filter(asset_status=Asset.Status.SCRAPPED)
    elif status_filter == 'archived':
        qs = qs.filter(asset_status=Asset.Status.ARCHIVED)

    if model_query:
        qs = qs.filter(
            Q(model__icontains=model_query) | Q(item_type__icontains=model_query) |
            Q(brand__icontains=model_query) | Q(computer_name__icontains=model_query) |
            Q(computer_name_new__icontains=model_query)
        )

    grouped = list(qs.annotate(
        model_label=Case(
            When(model__gt='', then=F('model')),
            When(item_type__gt='', then=F('item_type')),
            default=Value('غير محدد'), output_field=CharField(),
        ),
        brand_label=Case(
            When(brand__gt='', then=F('brand')),
            default=Value('غير محدد'), output_field=CharField(),
        ),
    ).values('brand_label', 'model_label').annotate(count=Count('id')).order_by('-count', 'brand_label', 'model_label'))
    return {
        'model_type': model_type,
        'status_filter': status_filter,
        'model_query': model_query,
        'model_rows': grouped,
        'model_total': sum(row['count'] for row in grouped),
    }


def _chart_rows(queryset, field_name, limit=8):
    rows = list(queryset.annotate(
        chart_label=Case(
            When(**{f'{field_name}__gt': ''}, then=F(field_name)),
            default=Value('غير محدد'), output_field=CharField(),
        )
    ).values('chart_label').annotate(count=Count('id')).order_by('-count', 'chart_label')[:limit])
    max_count = max([row['count'] for row in rows] or [1])
    for row in rows:
        row['percent'] = max(5, int(row['count'] * 100 / max_count))
    return rows


def _dashboard_context(request, paginate=True):
    from accounts.models import User

    tracked = _tracked_assets()
    hardware = _hardware_assets()
    assigned_qs = tracked.filter(asset_status=Asset.Status.ACTIVE).exclude(hr_code='')
    stock_qs = tracked.filter(Q(asset_status=Asset.Status.IN_STOCK) | Q(asset_status=Asset.Status.ACTIVE, hr_code=''))
    scrapped_qs = tracked.filter(asset_status=Asset.Status.SCRAPPED)
    pending_qs = tracked.filter(asset_status=Asset.Status.PENDING)
    archived_qs = tracked.filter(asset_status=Asset.Status.ARCHIVED)

    computer_qs = hardware.filter(asset_type=Asset.Types.COMPUTER)
    laptop_q = (
        Q(item_type__icontains='laptop') | Q(model__icontains='laptop') |
        Q(item_type__icontains='notebook') | Q(model__icontains='notebook')
    )
    model_info = _dashboard_model_rows(request)
    rows = model_info['model_rows']
    if paginate:
        models_page = Paginator(rows, 25).get_page(request.GET.get('page'))
    else:
        models_page = None

    type_data = []
    active_for_charts = tracked.exclude(asset_status__in=[Asset.Status.SCRAPPED, Asset.Status.ARCHIVED])
    active_hardware = hardware.exclude(asset_status__in=[Asset.Status.SCRAPPED, Asset.Status.ARCHIVED])
    active_computers = active_hardware.filter(asset_type=Asset.Types.COMPUTER)
    active_laptop_count = active_computers.filter(laptop_q).count()
    active_pc_count = active_computers.exclude(laptop_q).count()
    chart_values = [
        ('PC / Desktop', active_pc_count),
        ('Laptop', active_laptop_count),
        ('الشاشات', active_hardware.filter(asset_type=Asset.Types.MONITOR).count()),
        ('الطابعات', active_hardware.filter(asset_type=Asset.Types.PRINTER).count()),
        ('تراخيص GStar CAD', active_for_charts.filter(asset_type=Asset.Types.GSTAR).count()),
        ('أصول أخرى', active_hardware.filter(asset_type=Asset.Types.OTHER).count()),
    ]
    for label, count in chart_values:
        type_data.append({'label': label, 'count': count})
    type_max = max([row['count'] for row in type_data] or [1])
    for row in type_data:
        row['percent'] = max(3, int(row['count'] * 100 / type_max)) if row['count'] else 0

    projects = _chart_rows(active_for_charts, 'project')
    departments = _chart_rows(active_for_charts, 'department')

    scrap_source_rows = list(Asset.objects.filter(asset_type=Asset.Types.SCRAP).values('id', 'serial_number'))
    scrapped_serials = set(
        tracked.filter(asset_status=Asset.Status.SCRAPPED).exclude(serial_number='').values_list('serial_number', flat=True)
    )
    scrapped_serials_folded = {serial.casefold() for serial in scrapped_serials}
    standalone_scraps = set()
    for row in scrap_source_rows:
        serial = (row['serial_number'] or '').strip()
        if serial and serial.casefold() in scrapped_serials_folded:
            continue
        standalone_scraps.add(serial.casefold() if serial else f"row:{row['id']}")
    scrapped_count = scrapped_qs.count() + len(standalone_scraps)

    return {
        'device_count': Device.objects.count(),
        'asset_count': Asset.objects.count(),
        'employee_count': Employee.objects.count(),
        'user_count': User.objects.count(),
        'is_admin': request.user.is_admin_role,
        'tracked_asset_count': tracked.count(),
        'active_count': assigned_qs.count(),
        'stock_count': stock_qs.count(),
        'scrapped_count': scrapped_count,
        'pending_count': pending_qs.count(),
        'archived_count': archived_qs.count(),
        'pc_count': computer_qs.exclude(laptop_q).count(),
        'laptop_count': computer_qs.filter(laptop_q).count(),
        'monitor_count': hardware.filter(asset_type=Asset.Types.MONITOR).count(),
        'printer_count': hardware.filter(asset_type=Asset.Types.PRINTER).count(),
        'type_data': type_data,
        'project_data': projects,
        'department_data': departments,
        'model_type': model_info['model_type'],
        'status_filter': model_info['status_filter'],
        'model_query': model_info['model_query'],
        'model_total': model_info['model_total'],
        'model_rows': rows,
        'models_page': models_page,
    }


@login_required
def dashboard(request):
    context = {
        'device_count': Device.objects.count(),
        'asset_count': Asset.objects.count(),
        'employee_count': Employee.objects.count(),
        'is_admin': request.user.is_admin_role,
    }
    if request.user.is_admin_role:
        context['user_count'] = __import__('accounts.models', fromlist=['User']).User.objects.count()
    return render(request, 'devices/dashboard.html', context)


@admin_required
def inventory_dashboard(request):
    return render(request, 'inventory/admin_dashboard.html', _dashboard_context(request))


@admin_required
def inventory_dashboard_export(request):
    context = _dashboard_context(request, paginate=False)
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['Metric', 'Count'])
    summary_rows = [
        ('Total imported/source rows', context['asset_count']),
        ('Tracked assets excluding source-only rows', context['tracked_asset_count']),
        ('Active custody', context['active_count']),
        ('In stock', context['stock_count']),
        ('Scrapped assets (deduplicated against Scraped source rows)', context['scrapped_count']),
        ('Pending', context['pending_count']),
        ('Archived', context['archived_count']),
        ('PC', context['pc_count']), ('Laptop', context['laptop_count']),
        ('Monitors', context['monitor_count']), ('Printers', context['printer_count']),
    ]
    for row in summary_rows:
        summary.append(list(row))

    models_ws = workbook.create_sheet('Model Counts')
    models_ws.append(['Asset group', 'Brand', 'Model / Item Type', 'Count', 'Status filter', 'Search filter'])
    for row in context['model_rows']:
        models_ws.append([
            context['model_type'], row['brand_label'], row['model_label'], row['count'],
            context['status_filter'], context['model_query'],
        ])

    for sheet_name, key, label in [
        ('By Project', 'project_data', 'Project'), ('By Department', 'department_data', 'Department'),
        ('By Type', 'type_data', 'Asset Type'),
    ]:
        ws = workbook.create_sheet(sheet_name)
        ws.append([label, 'Count'])
        for row in context[key]:
            ws.append([row['label'] if 'label' in row else row['chart_label'], row['count']])

    for ws in workbook.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            width = min(max(max(len(str(cell.value or '')) for cell in col) + 2, 12), 48)
            ws.column_dimensions[letter].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    log_event(request.user, 'DASHBOARD_EXPORT', request, 'Exported filtered management dashboard and model counts')
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="SIAC_Management_Dashboard.xlsx"'
    return response


@admin_required
def movement_list(request):
    query = (request.GET.get('q') or '').strip()
    event_type = (request.GET.get('event_type') or '').strip()
    movements = AssetMovement.objects.select_related('asset', 'from_employee', 'to_employee', 'created_by')
    if query:
        movements = movements.filter(
            Q(asset__serial_number__icontains=query) | Q(asset__model__icontains=query) |
            Q(asset__computer_name__icontains=query) | Q(asset__computer_name_new__icontains=query) |
            Q(from_hr_code__icontains=query) | Q(to_hr_code__icontains=query) |
            Q(from_employee_name__icontains=query) | Q(to_employee_name__icontains=query) |
            Q(reason__icontains=query) | Q(reference_number__icontains=query)
        )
    if event_type in dict(AssetMovement.EventTypes.choices):
        movements = movements.filter(event_type=event_type)
    pager = Paginator(movements.order_by('-created_at', '-id'), 50)
    page = pager.get_page(request.GET.get('page'))
    return render(request, 'inventory/movement_list.html', {
        'page': page, 'query': query, 'selected_event_type': event_type,
        'event_types': AssetMovement.EventTypes.choices, 'movement_count': movements.count(),
    })


@admin_required
def movement_asset_search(request):
    query = (request.GET.get('q') or '').strip()
    physical_asset_types = [Asset.Types.COMPUTER, Asset.Types.PRINTER, Asset.Types.MONITOR, Asset.Types.GSTAR, Asset.Types.OTHER]
    assets = Asset.objects.filter(asset_type__in=physical_asset_types).exclude(
        asset_status__in=[Asset.Status.SCRAPPED, Asset.Status.ARCHIVED]
    ).select_related('employee')
    if query:
        assets = assets.filter(
            Q(serial_number__icontains=query) | Q(asset_tag__icontains=query) |
            Q(computer_name__icontains=query) | Q(computer_name_new__icontains=query) |
            Q(model__icontains=query) | Q(brand__icontains=query) | Q(hr_code__icontains=query)
        )
    else:
        assets = assets.none()
    pager = Paginator(assets.order_by('asset_type', 'serial_number', 'id'), 30)
    page = pager.get_page(request.GET.get('page'))
    return render(request, 'inventory/movement_asset_search.html', {
        'query': query, 'page': page, 'match_count': assets.count(),
    })


@admin_required
def movement_create(request, asset_pk):
    asset = get_object_or_404(Asset.objects.select_related('employee'), pk=asset_pk)
    if asset.asset_type in (Asset.Types.SCRAP, Asset.Types.NEW_2026, Asset.Types.REQUEST_2025):
        messages.error(request, 'السجل المحدد ليس أصلًا تشغيليًا قابلًا لحركة عهدة مستقلة.')
        return redirect('movement_asset_search')

    form = CustodyMovementForm(request.POST or None, asset=asset)
    if request.method == 'POST' and form.is_valid():
        try:
            with transaction.atomic():
                locked = Asset.objects.select_for_update().select_related('employee').get(pk=asset.pk)
                # Revalidate against the latest locked state so two admins cannot transfer the same old state.
                form = CustodyMovementForm(request.POST, asset=locked)
                if not form.is_valid():
                    return render(request, 'inventory/movement_form.html', {'asset': locked, 'form': form})

                old_hr = (locked.hr_code or '').strip()
                old_employee = locked.employee if locked.employee_id else (Employee.objects.filter(hr_code__iexact=old_hr).first() if old_hr else None)
                old_name = (old_employee.name if old_employee else locked.employee_name or old_hr or 'المخزون').strip()
                old_status = locked.asset_status
                action = form.cleaned_data['event_type']
                target = form.cleaned_data.get('target_employee')
                movement_date = form.cleaned_data['effective_date']
                reason = form.cleaned_data['reason'].strip()
                reference = form.cleaned_data.get('reference_number', '').strip()
                notes = form.cleaned_data.get('notes', '').strip()

                if action in (AssetMovement.EventTypes.ASSIGNED, AssetMovement.EventTypes.TRANSFERRED):
                    locked.hr_code = target.hr_code
                    locked.employee = target
                    locked.employee_name = target.name
                    if target.project:
                        locked.project = target.project
                    if target.sector:
                        locked.sector = target.sector
                    if target.department:
                        locked.department = target.department
                    locked.asset_status = Asset.Status.ACTIVE
                    to_employee, to_hr, to_name = target, target.hr_code, target.name
                elif action == AssetMovement.EventTypes.RETURNED:
                    locked.hr_code = ''
                    locked.employee = None
                    locked.employee_name = ''
                    locked.asset_status = Asset.Status.IN_STOCK
                    to_employee, to_hr, to_name = None, '', 'المخزون'
                else:  # SCRAPPED
                    locked.hr_code = ''
                    locked.employee = None
                    locked.employee_name = ''
                    locked.asset_status = Asset.Status.SCRAPPED
                    locked.scrap_date = movement_date.isoformat()
                    locked.scrap_number = reference
                    to_employee, to_hr, to_name = None, '', 'التخريد'

                locked.save()
                movement = AssetMovement.objects.create(
                    asset=locked,
                    event_type=action,
                    effective_date=movement_date,
                    from_employee=old_employee,
                    to_employee=to_employee,
                    from_hr_code=old_hr,
                    from_employee_name=old_name,
                    to_hr_code=to_hr,
                    to_employee_name=to_name,
                    from_status=old_status,
                    to_status=locked.asset_status,
                    reason=reason,
                    reference_number=reference,
                    notes=notes,
                    created_by=request.user,
                )
                action_log = {
                    AssetMovement.EventTypes.ASSIGNED: 'ASSET_ASSIGN',
                    AssetMovement.EventTypes.TRANSFERRED: 'ASSET_TRANSFER',
                    AssetMovement.EventTypes.RETURNED: 'ASSET_RETURN',
                    AssetMovement.EventTypes.SCRAPPED: 'ASSET_SCRAP',
                }[action]
                log_event(
                    request.user, action_log, request,
                    f'{movement.get_event_type_display()}: asset_id={locked.pk}, serial={locked.serial_number}, from={old_hr or "stock"}, to={to_hr or to_name}, date={movement_date.isoformat()}, reference={reference}',
                )
                messages.success(request, f'تم تسجيل الحركة بنجاح للأصل {locked.serial_number or locked.model or locked.pk}. تم حفظ الحركة في السجل التاريخي.')
                return redirect('movement_list')
        except Asset.DoesNotExist:
            messages.error(request, 'الأصل لم يعد موجودًا؛ حدّث الصفحة ثم حاول مرة أخرى.')
            return redirect('movement_asset_search')
    return render(request, 'inventory/movement_form.html', {'asset': asset, 'form': form})


@login_required
def search_device(request):
    # The supported search keys are deliberately limited to Serial Number and HR Code.
    serial = normalize_serial(request.GET.get('serial'))
    hr_code = normalize_serial(request.GET.get('hr_code'))
    mode = 'serial' if serial else ('hr' if hr_code else '')
    matched_assets = []
    custody_assets = []
    employee = None
    legacy_device = None
    matched_owners = []
    history_assets = []
    history_groups = []
    resolved_hr_code = hr_code
    found = False

    if serial:
        matched_assets = list(Asset.objects.filter(serial_number__iexact=serial).select_related('employee').order_by('source_sheet', 'id'))
        if matched_assets:
            history_assets = _omit_duplicate_scrap_rows(matched_assets)
            history_assets = list(_with_movement_history(Asset.objects.filter(pk__in=[a.pk for a in history_assets]).select_related('employee').order_by('source_sheet', 'id')))
        if not matched_assets:
            legacy_device = Device.objects.filter(serial_number__iexact=serial).first()
        codes = []
        for asset in matched_assets:
            code = (asset.hr_code or (asset.employee.hr_code if asset.employee_id else '')).strip()
            if code and code.casefold() not in [c.casefold() for c in codes]:
                codes.append(code)
        for code in codes:
            owner = Employee.objects.filter(hr_code__iexact=code).first()
            owner_assets = list(_custody_assets(code))
            sample = next((a for a in matched_assets if a.hr_code.casefold() == code.casefold()), None)
            matched_owners.append({
                'hr_code': code,
                'employee': owner,
                'employee_name': owner.name if owner else (sample.display_employee_name if sample else ''),
                'project': owner.project if owner else (sample.project if sample else ''),
                'asset_count': len(owner_assets),
            })
        if len(matched_owners) == 1:
            resolved_hr_code = matched_owners[0]['hr_code']
            employee = matched_owners[0]['employee']
            custody_assets = list(_custody_assets(resolved_hr_code))
        found = bool(matched_assets or legacy_device)
        log_event(request.user, 'ASSET_SEARCH', request, f'Searched serial {serial}; found={found}; matches={len(matched_assets)}')
    elif hr_code:
        employee = Employee.objects.filter(hr_code__iexact=hr_code).first()
        matched_assets = list(Asset.objects.filter(hr_code__iexact=hr_code).select_related('employee').order_by('asset_type', 'id'))
        custody_assets = list(_custody_assets(hr_code))
        movement_asset_ids = AssetMovement.objects.filter(
            Q(from_hr_code__iexact=hr_code) | Q(to_hr_code__iexact=hr_code)
        ).values_list('asset_id', flat=True).distinct()
        history_assets = list(_with_movement_history(Asset.objects.filter(
            Q(pk__in=movement_asset_ids) | Q(hr_code__iexact=hr_code)
        ).select_related('employee').distinct().order_by('asset_type', 'serial_number', 'id')))
        history_assets = _omit_duplicate_scrap_rows(history_assets)
        found = bool(employee or matched_assets or history_assets)
        log_event(request.user, 'ASSET_SEARCH', request, f'Searched HR Code {hr_code}; found={found}; assets={len(matched_assets)}; history={len(history_assets)}')

    if history_assets:
        history_groups = _movement_history_groups(history_assets)

    # Avoid template access such as matched_owners.0 when a direct HR-code
    # search has no serial-search owner list. Django resolves template filter
    # arguments eagerly, so `default:matched_owners.0...` can raise 500 even
    # when `employee.name` already has a value.
    primary_owner = matched_owners[0] if matched_owners else {}
    context = {
        'serial': serial,
        'hr_code': hr_code,
        'mode': mode,
        'found': found,
        'employee': employee,
        'resolved_hr_code': resolved_hr_code,
        'matched_assets': matched_assets,
        'custody_assets': custody_assets,
        'matched_owners': matched_owners,
        'primary_owner_name': primary_owner.get('employee_name', ''),
        'primary_owner_project': primary_owner.get('project', ''),
        'legacy_device': legacy_device,
        'active_count': len(custody_assets),
        'history_groups': history_groups,
    }
    return render(request, 'devices/search.html', context)


@login_required
def download_custody_report(request, hr_code):
    code = normalize_serial(hr_code)
    employee = Employee.objects.filter(hr_code__iexact=code).first()
    assets = list(_custody_assets(code))
    all_matching = Asset.objects.filter(hr_code__iexact=code).exclude(asset_type=Asset.Types.SCRAP).exists()
    if not employee and not all_matching:
        messages.error(request, 'لا يوجد موظف أو أصول مرتبطة بكود HR هذا.')
        return redirect('search_device')
    log_event(request.user, 'CUSTODY_PDF_DOWNLOAD', request, f'Downloaded custody PDF for HR Code {code}')
    filename = slugify(code) or 'employee'
    return FileResponse(
        build_custody_pdf(employee, assets, hr_code=code),
        as_attachment=True,
        filename=f'SIAC_Custody_{filename}.pdf',
        content_type='application/pdf',
    )


@login_required
def download_report(request, pk):
    device = get_object_or_404(Device, pk=pk)
    log_event(request.user, 'PDF_DOWNLOAD', request, f'Downloaded legacy device report for {device.serial_number}')
    return FileResponse(build_pdf(device), as_attachment=True, filename=f'Report_{device.serial_number}.pdf', content_type='application/pdf')


# ---- Admin-only employee and asset management ----
@admin_required
def employee_list(request):
    query = (request.GET.get('q') or '').strip()
    employees = Employee.objects.all()
    if query:
        employees = employees.filter(
            Q(hr_code__icontains=query) | Q(name__icontains=query) | Q(arabic_name__icontains=query)
            | Q(project__icontains=query) | Q(department__icontains=query)
        )
    pager = Paginator(employees, 50)
    page = pager.get_page(request.GET.get('page'))
    return render(request, 'inventory/employee_list.html', {'page': page, 'query': query, 'employee_count': employees.count()})


@admin_required
def employee_create(request):
    form = EmployeeForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        employee = form.save()
        log_event(request.user, 'EMPLOYEE_CREATE', request, f'Created employee HR Code {employee.hr_code}')
        messages.success(request, f'تمت إضافة الموظف {employee.name}.')
        return redirect('employee_list')
    return render(request, 'inventory/employee_form.html', {'form': form, 'title': 'إضافة موظف جديد'})


@admin_required
def employee_edit(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    old_hr_code = employee.hr_code
    form = EmployeeForm(request.POST or None, instance=employee)
    if request.method == 'POST' and form.is_valid():
        updated = form.save()
        if old_hr_code != updated.hr_code:
            Asset.objects.filter(hr_code=old_hr_code).update(hr_code=updated.hr_code, employee=updated)
        log_event(request.user, 'EMPLOYEE_EDIT', request, f'Edited employee HR Code {updated.hr_code}')
        messages.success(request, 'تم حفظ بيانات الموظف وتحديث روابط عهدته.')
        return redirect('employee_list')
    return render(request, 'inventory/employee_form.html', {'form': form, 'title': f'تعديل الموظف: {employee.name}'})


@admin_required
def asset_list(request):
    query = (request.GET.get('q') or '').strip()
    asset_type = (request.GET.get('type') or '').strip()
    assets = Asset.objects.select_related('employee').all()
    if query:
        assets = assets.filter(
            Q(serial_number__icontains=query) | Q(hr_code__icontains=query)
            | Q(employee_name__icontains=query) | Q(employee__name__icontains=query)
            | Q(computer_name_new__icontains=query) | Q(computer_name__icontains=query)
            | Q(model__icontains=query) | Q(project__icontains=query)
        )
    if asset_type in dict(Asset.Types.choices):
        assets = assets.filter(asset_type=asset_type)
    pager = Paginator(assets.order_by('asset_type', 'serial_number', 'id'), 50)
    page = pager.get_page(request.GET.get('page'))
    return render(request, 'inventory/asset_list.html', {
        'page': page, 'query': query, 'selected_type': asset_type,
        'asset_types': Asset.Types.choices, 'asset_count': assets.count(),
    })


@admin_required
def asset_create(request):
    form = AssetForm(request.POST or None)
    if request.method == 'GET':
        # New physical assets start in stock; assigning them is a separate history event.
        form.initial['asset_status'] = Asset.Status.IN_STOCK
    if request.method == 'POST' and form.is_valid():
        asset = form.save()
        _record_asset_movement(request, asset, form, before=None)
        log_event(request.user, 'ASSET_CREATE', request, f'Created asset type={asset.asset_type}, serial={asset.serial_number}')
        if asset.hr_code and not asset.employee_id:
            messages.warning(request, 'تم حفظ الأصل، لكن HR Code غير موجود في قائمة الموظفين؛ ستظل البيانات محفوظة ويمكن ربطه لاحقًا.')
        else:
            messages.success(request, 'تمت إضافة الأصل بنجاح.')
        return redirect('asset_list')
    return render(request, 'inventory/asset_form.html', {'form': form, 'title': 'إضافة أصل جديد'})


@admin_required
def asset_edit(request, pk):
    asset = get_object_or_404(Asset.objects.select_related('employee'), pk=pk)
    before = Asset.objects.select_related('employee').get(pk=asset.pk)
    form = AssetForm(request.POST or None, instance=asset)
    for protected_field in ('hr_code', 'asset_status', 'scrap_date', 'scrap_number'):
        if protected_field in form.fields:
            form.fields[protected_field].disabled = True
    form.fields.pop('movement_date', None)
    form.fields.pop('movement_reason', None)
    form.fields['hr_code'].help_text = 'تغيير الموظف يتم من شاشة حركات العهدة لضمان حفظ التاريخ.'
    form.fields['asset_status'].help_text = 'تغيير الحالة يتم من شاشة حركات العهدة لضمان حفظ التاريخ.'
    if request.method == 'POST' and form.is_valid():
        updated = form.save()
        _record_asset_movement(request, updated, form, before=before)
        log_event(request.user, 'ASSET_EDIT', request, f'Edited asset id={updated.pk}, serial={updated.serial_number}')
        if updated.hr_code and not updated.employee_id:
            messages.warning(request, 'تم حفظ التعديل، لكن HR Code غير موجود في قائمة الموظفين.')
        else:
            messages.success(request, 'تم حفظ تعديلات الأصل.')
        return redirect('asset_list')
    return render(request, 'inventory/asset_form.html', {'form': form, 'title': f'تعديل الأصل: {asset.serial_number or asset.model or asset.pk}'})


@admin_required
def inventory_import(request):
    form = InventoryWorkbookImportForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and form.is_valid():
        workbook = form.cleaned_data['workbook']
        try:
            summary = preview_inventory_workbook(workbook)
            pending = PendingInventoryImport.objects.create(
                file=workbook, uploaded_by=request.user, summary=summary
            )
            log_event(request.user, 'INVENTORY_PREVIEW', request, f"Previewed inventory workbook: employees={summary['employees']}, assets={summary['total_assets']}")
            return redirect('inventory_import_preview', pk=pending.pk)
        except Exception as exc:
            messages.error(request, f'تعذر قراءة الملف: {exc}')
    return render(request, 'inventory/import.html', {'form': form})


@admin_required
def inventory_import_preview(request, pk):
    pending = get_object_or_404(PendingInventoryImport, pk=pk, uploaded_by=request.user, is_approved=False)
    if request.method == 'POST':
        try:
            with pending.file.open('rb') as workbook:
                result = import_inventory_workbook(workbook, actor=request.user)
            pending.is_approved = True
            pending.save(update_fields=['is_approved'])
            details = ', '.join(f'{name}: {count}' for name, count in result['by_sheet'].items())
            log_event(request.user, 'INVENTORY_IMPORT', request, f"Imported inventory workbook: employees={result['employees']}, assets={result['assets']}; {details}")
            messages.success(request, f"اكتمل الاستيراد: {result['employees']} موظفًا و{result['assets']} سجل أصل. تم تحديث السجلات المستوردة دون حذف الإدخالات اليدوية.")
            return redirect('asset_list')
        except Exception as exc:
            messages.error(request, f'تعذر استيراد الملف: {exc}')
    return render(request, 'inventory/import_preview.html', {'pending': pending, 'summary': pending.summary})


@admin_required
def inventory_export(request):
    response = build_inventory_excel_response()
    log_event(request.user, 'INVENTORY_EXPORT', request, f'Exported inventory workbook: assets={Asset.objects.count()}, employees={Employee.objects.count()}')
    return response


# Legacy single-sheet Excel import remains available for compatibility with the initial version.
@admin_required
def manage_excel(request):
    current = DeviceImport.objects.filter(is_current=True).first()
    form = ExcelUploadForm(request.POST or None, request.FILES or None)
    if request.method == 'POST' and form.is_valid():
        try:
            pending = create_pending_import(form.cleaned_data['excel_file'], request.user)
            log_event(request.user, 'EXCEL_PREVIEW', request, f'Prepared legacy Excel preview: {pending.valid_rows} valid, {pending.duplicate_rows} duplicates')
            return redirect('excel_preview', pk=pending.pk)
        except Exception as exc:
            messages.error(request, f'فشل قراءة الملف: {exc}')
    return render(request, 'devices/manage_excel.html', {'form': form, 'current': current})


@admin_required
def excel_preview(request, pk):
    pending = get_object_or_404(PendingExcelImport, pk=pk, uploaded_by=request.user, is_approved=False)
    if request.method == 'POST':
        try:
            with pending.file.open('rb') as f:
                count, record, info = import_excel(f, request.user)
            pending.is_approved = True
            pending.save(update_fields=['is_approved'])
            log_event(request.user, 'EXCEL_IMPORT', request, f'Approved legacy Excel import: {count} devices; duplicates ignored={info["duplicate_rows"]}')
            messages.success(request, f'تم اعتماد الملف بنجاح. تم تحميل {count} جهاز إلى قاعدة البيانات القديمة.')
            return redirect('manage_excel')
        except Exception as exc:
            messages.error(request, f'فشل اعتماد الملف: {exc}')
    return render(request, 'devices/excel_preview.html', {'pending': pending, 'current_count': Device.objects.count()})
